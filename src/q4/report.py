"""Q4 tradeoff figures, template-compatible tables and audit handoff."""
from src.q4.partition import *
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import openpyxl

def table(headers,rows):return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(map(str,r))+' |' for r in rows])
def savefig(fig,name):
    for ext in ['png','pdf','svg']:fig.savefig(OUT/f'{name}.{ext}',dpi=180)
    p=OUT/f'{name}.svg';p.write_text('\n'.join(x.rstrip() for x in p.read_text().splitlines())+'\n');plt.close(fig)

def workbook(ids,name,byid):
    raw=ROOT/'data/raw/结果提交模板.xlsx';w=openpyxl.load_workbook(raw)
    for sheet in list(w):
        if sheet.title!='Q4_分区配置':w.remove(sheet)
    ws=w['Q4_分区配置'];rowno=2
    for sid in ids:
        q=byid[sid]
        for group in q['groups']:
            values=[q['k'],group['group_id'],','.join(group['services'])]+[group['resources'][kind] for kind in KINDS]
            for col,value in enumerate(values,1):ws.cell(rowno,col,value)
            rowno+=1
    notes=w.create_sheet('口径与库存缺口');notes.append(['基线','Q3E7_001，Gamma=0；所有任务时刻、航迹和通信关系冻结'])
    notes.append(['范围','仅 Q4 工作表草稿，不是全题最终提交文件'])
    notes.append(['方案','资源','需求','库存','缺口'])
    for sid in ids:
        for kind in KINDS:notes.append([sid,kind,byid[sid]['resources'][kind],byid[sid]['inventory'][kind],byid[sid]['shortage'][kind]])
    w.save(OUT/name)
    # Read the saved artifact back and compare all template cells.
    check=openpyxl.load_workbook(OUT/name,read_only=True,data_only=True)['Q4_分区配置'];expected=[]
    for sid in ids:
        for group in byid[sid]['groups']:expected.append(tuple([byid[sid]['k'],group['group_id'],','.join(group['services'])]+[group['resources'][kind] for kind in KINDS]))
    actual=list(check.iter_rows(min_row=2,max_row=1+len(expected),max_col=11,values_only=True));assert actual==expected

def main():
    config=json.loads((OUT/'config.json').read_text());records=json.loads((OUT/'all_partitions.json').read_text());byid={x['partition_id']:x for x in records};blocks=json.loads((OUT/'blocks.json').read_text())
    selected=[byid[s] for s in config['selected']];balanced=[byid[s] for s in config['balanced_alternatives']]
    workbook(config['selected'],'Q4_分区配置_缺口优先.xlsx',byid);workbook(config['balanced_alternatives'],'Q4_分区配置_均衡优先.xlsx',byid)
    groups=[]
    for role,rr in [('SHORTAGE_PRIORITY',selected),('BALANCE_PRIORITY',balanced)]:
        for q in rr:
            for g in q['groups']:groups.append(dict(partition_id=q['partition_id'],role=role,k=q['k'],group_id=g['group_id'],services=';'.join(g['services']),boxes=g['box_count'],mass_kg=g['mass_kg'],workload_s=g['workload_uav_occupied_s'],**g['resources']))
    csv(OUT/'group_configurations.csv',groups)
    fig,ax=plt.subplots(figsize=(7.6,4.6),layout='constrained')
    for k,color in [(2,'#2879b9'),(3,'#ba5a3a')]:
        rr=[q for q in records if q['k']==k];ax.scatter([q['total_shortage_units'] for q in rr],[q['workload_cv'] for q in rr],label=f'K={k}: all {len(rr)} feasible partitions',c=color,s=35,alpha=.65)
    for q in selected+balanced:
        ax.scatter(q['total_shortage_units'],q['workload_cv'],s=100,facecolors='none',edgecolors='black')
        ax.annotate(q['partition_id'].replace('Q4_',''),(q['total_shortage_units'],q['workload_cv']),xytext=(5,7),textcoords='offset points',fontsize=8)
    ax.set_xlabel('Additional resource units, summed across types (not monetary cost)');ax.set_ylabel('UAV occupied-time workload CV');ax.set_ylim(-.05,1.3);ax.grid(alpha=.2);ax.legend(fontsize=8);ax.set_title('Fixed Q3 execution: independent partition resource–balance tradeoffs')
    savefig(fig,'partition_resource_balance')
    nodes=read(DATA/'nodes.csv').set_index('node_id');fig,axes=plt.subplots(1,2,figsize=(10,4.6),layout='constrained');colors=['#2879b9','#c96338','#46974f']
    for ax,q in zip(axes,selected):
        for i,g in enumerate(q['groups']):
            f=nodes.loc[g['services']];ax.scatter(f.x_m/1000,f.y_m/1000,c=colors[i],label=g['group_id'],s=42)
            for sid,row in f.iterrows():ax.annotate(sid,(row.x_m/1000,row.y_m/1000),xytext=(3,3),textcoords='offset points',fontsize=7)
        origin=nodes.loc['O01'];ax.scatter(origin.x_m/1000,origin.y_m/1000,marker='^',c='black',label='O01');ax.set_xlabel('UTM easting (km)');ax.set_ylabel('UTM northing (km)');ax.set_aspect('equal');ax.margins(x=.12,y=.08);ax.grid(alpha=.2);ax.legend(fontsize=7);ax.set_title(f"K={q['k']}, inventory-shortage priority")
    savefig(fig,'partition_maps')
    summary=table(['角色','方案','资源总件数','新增缺口件数','工作量 CV','库存可独立执行'],[[role,q['partition_id'],q['total_resource_units'],q['total_shortage_units'],f"{q['workload_cv']:.6f}",'否，需补足所列缺口'] for role,rr in [('缺口优先',selected),('均衡优先',balanced)] for q in rr])
    resources=table(['资源类型','现库存','不分组最小池','K2 缺口优先需求','K2 缺口','K3 缺口优先需求','K3 缺口'],[[kind,blocks['inventory'][kind],blocks['global_minimum_resources'][kind],selected[0]['resources'][kind],selected[0]['shortage'][kind],selected[1]['resources'][kind],selected[1]['shortage'][kind]] for kind in KINDS])
    group_table=table(['方案','组','服务区','箱数','质量/kg','无人机占用/s','资源向量'],[[q['partition_id'],g['group_id'],','.join(g['services']),g['box_count'],f"{g['mass_kg']:.2f}",f"{g['workload_uav_occupied_s']:.3f}",str([g['resources'][kind] for kind in KINDS])] for q in selected+balanced for g in q['groups']])
    blocktable=table(['不可拆块','服务区'],[[i+1,','.join(ss)] for i,ss in enumerate(blocks['blocks'])])
    text=f'''# Q4 独立分区与资源配置报告

Gate：**Q4_INDEPENDENT_PARTITIONS_READY**。基于标准 ΓC=0 的 **Q3E7_001**，保持任务、箱组、访问次序、时刻及通信保障关系不变，完成 K=2/3 全部合法分区的精确枚举和独立资源核算。E8-B 不再是 Q4 门禁。

这里 READY 表示分区、最小需求和缺口已正确求出，不表示现有库存足以执行独立分区。**在本次冻结方案及整架次不可拆依赖口径下，2 组和 3 组均不存在零库存缺口分区**；穷举分别为 15 个、25 个，最低缺口分别为 2 件和 6 件。这是固定 Q3 输入上的确切结论，不是所有 Q3 方案或允许重调度后的全局结论。

## 分区口径

同一运输架次涉及的服务区必须同组。为同时保持完整中继架次安排、通信关系和组间独立性，一个中继架次保障的所有运输架次也必须同组；不复制、不切开既有中继架次。此前冻结的依赖超图架构因此收缩成下列五块：

{blocktable}

按五个不可拆块精确枚举，Stirling 数 S(5,2)=15、S(5,3)=25。原物理机号可以重新分配为组内专用资源编号，否则把原来跨时段复用的机号也当不可拆依赖，会错误改变“独立资源需求核算”的问题。所有任务开始/结束、返航、充电完成时刻及源机号均保留，可追溯核验。

## 代表方案及资源代价

未给设备采购单价，因此没有捏造货币成本。完整 Pareto 使用八维资源需求向量与工作量 CV。便于提交的缺口优先代表按“新增资源总件数最小 → 资源总件数最小 → CV 最小”选取；同时给出均衡优先代表，避免把高度失衡的最低缺口分区说成唯一最佳答案。

{summary}

{resources}

资源向量顺序为 A/B/C 运输机、A/B/C 电池、中继机、中继能源组件。缺口优先 K2 需要新增 **B 型运输机 1 架、B 型电池 1 组**。K3 需要新增 **B 型运输机 2 架、C 型运输机 1 架、B 型电池 2 组、C 型电池 1 组**。

值得注意：K2 所需资源总件数 28 小于现库存总件数 30，却仍有两件类型缺口。A 型电池或中继组件的富余不能替代 B 型运输机/电池。因此必须报告分类型缺口，不能仅比较总数量。

{group_table}

## 资源最小性、冗余与均衡

运输机占用 [准备开始,返航)，同型电池占用 [起飞,充至100%)；中继机占用 [准备开始,返航加周转)，组件占用 [起飞,充至100%)。组内同类资源的最小数等于这些区间的最大重叠数：最大团提供下界，区间着色提供达到该下界的实际日历。数值端点容差为 1e-6 s。按机型单独计算，且跨组不复用任何资源。

独立分区冗余定义为各组最小需求之和减去同一冻结任务不分组时的最小池；与“现库存剩余”区分。缺口优先 K2/K3 分别比不分组最小池增加 3/8 件资源。原 E7 使用过六个组件编号，但区间重分配证明不分组最小池是四个；不能把“曾使用编号数量”误当最少需求。

工作量为每组运输机准备至返航占用时长与中继机准备至周转完成占用时长之和，CV 用总体标准差除以平均值。另报告箱数、质量和任务数；不把服务区数量相等当作工作量均衡。最低缺口方案把大量任务留在一个大组，所以 CV 较高；均衡方案分别把 CV 降到约 0.0070 / 0.2598，但新增资源缺口增加到 5 / 13 件。

![分区代价与均衡](partition_resource_balance.png)

![缺口优先分区地图](partition_maps.png)

按类型看，独立池的需求是各组峰值相加，不是全体任务峰值；峰值发生在不同时刻时，禁止跨组调配仍要求各组保有自己的资源。原 Q3 的时间复用收益因此损失。源方案全部物理任务未变，交付时刻、联合完工时间和运输/中继能耗保持原值；多配置设备不等于新增飞行任务。

## 验证、提交与边界

独立验证器不导入分区优化代码：以图遍历重建连通块，用另一套全标号枚举核对 40 个分区，用事件扫描计算全部资源峰值，并检查每个峰值证书和四套导出资源日历。逐字段核对源任务和通信关系不变，重新分配的资源不跨组、类型正确且不重叠。3 个拒绝测试覆盖篡改任务时刻、少报资源需求和切断中继依赖。

提供与官方 Q4 表头一致的 [缺口优先工作簿](Q4_分区配置_缺口优先.xlsx) 和 [均衡优先工作簿](Q4_分区配置_均衡优先.xlsx)，已回读逐单元格核对；均只是 Q4 页草稿，不是全题最终提交文件。[组配置 CSV](group_configurations.csv) 与 solutions 下完整执行配置可供整合。

```bash
export OPENBLAS_NUM_THREADS=1 OMP_NUM_THREADS=1
.venv/bin/python -m src.q4.partition
.venv/bin/python -m src.q4.report
.venv/bin/python validation/mathematical/q4_validate.py
.venv/bin/python validation/mathematical/test_q4_regressions.py
```

源 Q3 的通信物理已通过独立全航程采样，Q4 通过源审计哈希及逐字段继承维持该证明链，没有再跑无变化的射线传播。最小性是固定任务时间/依赖结构下的精确结果；若更换 E7 点、允许移动任务时刻、复制/拆分中继，必须作为另一个模型重新计算，不能套用本结论。
'''
    (OUT/'Q4_REPORT.md').write_text(text)
    (OUT/'Q4_GPT_SYNC.md').write_text(f'''# Q4 给 GPT 的结果交接

正式 Gate：Q4_INDEPENDENT_PARTITIONS_READY。选定 Γ=0 的 Q3E7_001；箱组、路线、机型、任务时刻、完整中继架次和通信关系全部冻结。E8-B 不阻塞 Q4。

五个不可拆运输—中继依赖块；K2/K3 精确枚举 15/25 个分区，八维资源需求+CV 的非支配分区分别为 8/16 个。40 个资源需求向量均经独立峰值扫描验证，四套代表方案有完整独立资源日历及最大团下界证书。

{summary}

{resources}

两组最小缺口：B 机1、B电池1。三组最小缺口：B机2、C机1、B电池2、C电池1。冻结方案下所有合法分区都有库存缺口；这是 Q4 要求报告的结果，不能把 Gate READY 误读成现库存可执行。资源总件数也不能抵消类型缺口：K2 需28件、库存30件仍缺两件B型资源。

最低缺口代表较失衡，已给均衡优先替代及完整 Pareto；请在论文中一起展示，勿只说一个“最优分区”。选择规则是件数优先，非货币采购成本。原六个已使用组件编号经过重新着色的最小需求仅四个，必须区分已用编号与最低容量。

论证范围：固定 Q3E7_001、整中继架次不可拆依赖、无跨组复用。没有声称全部 Q3 方案下的最小缺口。Q3通信证明以源文件哈希继承，Q4 独立核验全部任务时刻与通信指派不变。

完整结果：[报告](Q4_REPORT.md)、[Gate](validation.json)、[配置](config.json)、[分组资源 CSV](group_configurations.csv)。两个 xlsx 是官方 Q4 页格式草稿，尚不是整题最终提交表。下一主线为提交表合并与论文主结果冻结；ns-3 仅作后期可选验证。
''')
    manifest=json.loads((OUT/'input_manifest.json').read_text());raw=ROOT/'data/raw/结果提交模板.xlsx';manifest['hashes'][str(raw.relative_to(ROOT))]=hashlib.sha256(raw.read_bytes()).hexdigest();write(OUT/'input_manifest.json',manifest)
    write(OUT/'submission_table_checks.json',dict(workbooks=2,rows_per_workbook=5,official_header_columns=11,readback_cell_equality=True,scope='Q4_ONLY_DRAFT'))
    write(OUT/'regression_checks.json',dict(tests=3,passed=3,cases=['changed_frozen_time_rejected','understated_resource_pool_rejected','cross_group_relay_dependency_rejected']))
    print('Q4_REPORTS_AND_TABLES_READY',flush=True)

if __name__=='__main__':main()

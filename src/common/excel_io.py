"""Excel reading helpers (raw files are opened read-only)."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd


def load_sheet_frame(path: Path, sheet: str, header_row: int) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(c).strip() if c is not None else f"col{i}" for i, c in enumerate(rows[header_row - 1])]
    data = rows[header_row:]
    # drop fully empty rows
    data = [r for r in data if any(v is not None and str(v).strip() != "" for v in r)]
    return pd.DataFrame(data, columns=header)


def sheet_values(path: Path, sheet: str) -> list[tuple]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows

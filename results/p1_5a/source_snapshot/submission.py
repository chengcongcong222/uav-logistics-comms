"""One F13 clock throughout the official workbook and supplements."""
import ast,math
import pandas as pd
import openpyxl
from pyproj import Transformer
from src.p1_5a.common import *
KINDS=['TUAV_A','TUAV_B','TUAV_C','TBAT_A','TBAT_B','TBAT_C','RUAV','REC']
DEST=OUT/'submission'
BOOK=DEST/'结果提交表_F13_自主主例.xlsx'
def frame(name):return pd.read_csv(PACKAGE/f'{name}_F13.csv',float_precision='round_trip')
def rows_q2(t):return [[r.sortie_id,r.uav_id,r.uav_type,r.battery_id,r.preparation_start_s,r.service_sequence,r.return_s,r.energy_kwh] for r in t.itertuples()]
def rows_delivery(b):return [[r.box_id,r.sortie_id,r.service_id,r.delivery_time_s] for r in b.itertuples()]
def communication(t,g):
    rows=[]
    for r in t.itertuples():
        gg=g[g.transport_sortie_id==r.sortie_id]
        cuts=sorted(set([r.takeoff_s,r.return_s]+gg.service_start_s.tolist()+gg.service_end_s.tolist()))
        for a,b in zip(cuts[:-1],cuts[1:]):
            if b-a<1e-8:continue
            active=gg[(gg.service_start_s<=(a+b)/2)&(gg.service_end_s>=(a+b)/2)]
            assert len(active)<=1
            rows.append([r.sortie_id,'全航程保障区间',a,b,'TRANSPORT_RELAY_G01' if len(active) else 'TRANSPORT_G01',active.iloc[0].relay_sortie_id if len(active) else None])
    return rows
def clean(v):return None if pd.isna(v) else v
def addframe(w,name,f):
    s=w.create_sheet(name);s.append(list(f.columns))
    for row in f.itertuples(index=False,name=None):s.append([clean(v) for v in row])
def build():
    require_admitted()
    assert read(OUT/'q4/independent_validation.json')['status']=='PASSED'
    DEST.mkdir(exist_ok=True)
    w=openpyxl.load_workbook(ROOT/'data/raw/结果提交模板.xlsx')
    official_headers={}
    for s in w:
        header=[v for v in next(s.values) if v is not None]
        official_headers[s.title]=header
        s.delete_rows(1,s.max_row);s.append(header)
    t=frame('transport_sorties');b=frame('box_delivery');r=frame('relay_sorties');g=frame('communication_guarantee')
    pk=pd.read_csv(ROOT/'results/reset/q1/q1_packings_rho20.csv',float_precision='round_trip')
    types=pd.read_csv(GEOMETRY/'transport_uav_types.csv').set_index('uav_type')
    data={'Q1_单点组批':[[f'Q1_{i:03d}',x.service_id,x.uav_type,','.join(ast.literal_eval(x.box_ids)),x.mass_kg,x.volume_m3,x.time_s,x.energy_kwh,100*(1-x.energy_kwh/types.loc[x.uav_type].battery_energy_kwh)] for i,x in enumerate(pk.itertuples(),1)],
      'Q2_运输架次':rows_q2(t),'Q2_逐箱交付':rows_delivery(b),'Q3_中继架次':[],
      'Q3_通信保障':communication(t,g),'Q4_分区配置':[]}
    sites=frame('candidate_sites').set_index('site_id');tf=Transformer.from_crs(32649,4326,always_xy=True)
    for x in r.itertuples():
        s=sites.loc[x.site_id];lon,lat=tf.transform(s.x_m,s.y_m)
        data['Q3_中继架次'].append([x.relay_sortie_id,x.relay_id,x.energy_component_id,x.preparation_start_s,lon,lat,s.z_amsl_m,x.service_start_s,x.service_end_s,x.return_s,x.total_energy_kwh])
    config=read(OUT/'q4/config.json');alternatives=[];q4rows=[];calendar=[];service=[];work=[]
    for sid in dict.fromkeys(config['selected']+config['balanced_alternatives']):
        q=read(OUT/'q4/solutions'/sid/'configuration.json')
        role='缺口优先' if sid in config['selected'] else '均衡优先'
        for group in q['groups']:
            row=[q['k'],group['group_id'],','.join(group['services'])]+[group['resources'][k] for k in KINDS]
            if role=='缺口优先':data['Q4_分区配置'].append(row)
            alternatives.append([sid,role]+row)
            for s in group['services']:service.append(dict(partition_id=sid,role=role,k=q['k'],service_id=s,group_id=group['group_id']))
            work.append(dict(partition_id=sid,role=role,k=q['k'],group_id=group['group_id'],box_count=group['box_count'],
                mass_kg=group['mass_kg'],transport_sorties=group['transport_sorties'],relay_sorties=group['relay_sorties'],
                workload_uav_occupied_s=group['workload_uav_occupied_s'],workload_cv=q['workload_cv'],
                occupied_resource_time_s=group['occupied_resource_time_s'],allocated_resource_time_s=group['allocated_resource_time_s']))
        for k in KINDS:q4rows.append(dict(partition_id=sid,role=role,k=q['k'],resource_type=k,required=q['resources'][k],
          inventory=q['inventory'][k],shortage=q['shortage'][k],inventory_surplus=q['inventory_surplus'][k],partition_redundancy=q['redundancy'][k]))
        calendar += [dict(partition_id=sid,role=role,**x) for x in q['calendar']]
    for name,rows in data.items():
        for row in rows:w[name].append(row)
    for name,header,rows in [
      ('Q3_运输时刻_必读',official_headers['Q2_运输架次'],rows_q2(t)),
      ('Q3_逐箱交付_必读',official_headers['Q2_逐箱交付'],rows_delivery(b)),
      ('Q4_双准则代表',['方案','选择规则']+official_headers['Q4_分区配置'],alternatives)]:
        s=w.create_sheet(name);s.append(header)
        for row in rows:s.append(row)
    for name,stem in [('运输完整时刻','transport_sorties'),('运输机资源日历','transport_uav_calendar'),
      ('运输电池资源日历','transport_battery_calendar'),('中继完整时刻','relay_sorties'),
      ('中继机资源日历','relay_uav_calendar'),('中继能源资源日历','relay_energy_calendar'),('原子通信保障关系','communication_guarantee')]:
        addframe(w,name,frame(stem))
    for name,rows in [('Q4_分类型缺口冗余',q4rows),('Q4_工作量与均衡',work),('Q4_分区归属',service),('Q4_独立资源日历',calendar)]:
        addframe(w,name,pd.DataFrame(rows))
    m=read(PACKAGE/'joint_metrics_F13.json');q1=read(ROOT/'results/reset/q1/validation.json')
    indicators=[['Q1箱数',80,'箱'],['Q1架次',len(pk),'架次'],['Q1能耗',float(pk.energy_kwh.sum()),'kWh'],
       ['Q2/Q3运输架次',len(t),'架次'],['Q2/Q3运输能耗',float(t.energy_kwh.sum()),'kWh'],
       ['Q2/Q3运输完工',float(t.return_s.max()),'s'],['J_late',m['J_late'],'加权秒'],
       ['J_norm',m['J_norm'],'无量纲'],['Q3联合完工',m['joint_makespan_s'],'s'],
       ['Q3联合完工分钟',m['joint_makespan_s']/60,'min'],['Q3总能耗',m['total_energy_kwh'],'kWh'],
       ['Q3中继能耗',float(r.total_energy_kwh.sum()),'kWh'],['Q3中继架次',len(r),'架次'],['Gamma_C',0,'dB']]
    s=w.create_sheet('主指标核对');s.append(['指标','数值','单位'])
    for row in indicators:s.append(row)
    s=w.create_sheet('口径与来源');s.append(['项目','说明'])
    notes=[
      ['主例',IDENT],['准入状态','CURRENT_AUTONOMOUS_MAIN_CANDIDATE'],
      ['来源','官方原始数据、自主2903模式池、重新独立验证的F13完整执行包'],
      ['Q1','G2原生DEM穿格几何；rho=0.20自主单点组批结果'],
      ['Q2/Q3同一运输时钟','Q2两页与Q3运输补充页完全相同，均采用F13最终联合调度的实际时刻；Q2表示该主例的运输层可行解'],
      ['Q2指标范围','运输能耗与运输完工仅计运输；Q3总能耗与联合完工包含中继'],
      ['时间与能量单位','时刻、时长均为秒；能耗kWh；主指标另列分钟；SOC百分比'],
      ['准备开始','官方运输及中继开始时刻为准备开始；完整表列起飞、返航、周转和充电完成'],
      ['通信验证','0.5秒全航程数值采样，包含阶段及保障边界，边界细化至0.1秒；不是解析连续时间证明'],
      ['Q4固定域','固定F13箱组、访问顺序、运输与中继时刻和保障关系；同一中继架次关联服务区归于同组'],
      ['Q4独立资源','各组独立配置，资源ID带组前缀，不跨组共享；缺口需要补足后才能按固定时刻独立执行'],
      ['Q4缺口优先','先最小资源缺口总件数，再资源配置总件数，再工作量CV；件数不等于货币成本'],
      ['Q4均衡优先','最小无人机占用工作量CV，缺口与配置规模另列'],
      ['资源冗余','分组需求合计减不分组的最低峰值需求；库存剩余另列，二者不混用'],
      ['适用边界','固定主例的可执行见证与任务分区；不声称全局最优或完整Pareto前沿']]
    for row in notes:s.append(row)
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for cell in s[1]:cell.font=openpyxl.styles.Font(bold=True)
    w.save(BOOK)
    write(DEST/'cell_payload.json',data)
    write(DEST/'input_manifest.json',dict(source_solution=IDENT,Q2_clock='F13_FINAL_Q3_TRANSPORT_CLOCK',
        Q4=config['selected'],balanced=config['balanced_alternatives'],
        hashes=hashes([ROOT/'data/raw/结果提交模板.xlsx',ROOT/'results/reset/q1/q1_packings_rho20.csv',
            *PACKAGE.glob('*.csv'),PACKAGE/'joint_metrics_F13.json',*[(OUT/'q4/solutions'/sid/'configuration.json') for sid in dict.fromkeys(config['selected']+config['balanced_alternatives'])]])))
    return data,official_headers

def check():
    """Re-read saved cells; compare every table plus derive global identities."""
    w=openpyxl.load_workbook(BOOK,data_only=False);checks=[]
    def rows(name,n=None):
        vals=list(w[name].values)[1:]
        return [list(x[:n]) if n else list(x) for x in vals if any(v is not None for v in x)]
    def equal(actual,expected,label):
        assert len(actual)==len(expected),(label,len(actual),len(expected))
        for i,(a,b) in enumerate(zip(actual,expected)):
            assert len(a)==len(b),(label,i)
            for j,(x,y) in enumerate(zip(a,b)):
                if isinstance(y,(int,float)) and not isinstance(y,bool):
                    assert x is not None and abs(x-y)<1e-8,(label,i,j,x,y)
                else:assert x==y,(label,i,j)
        checks.append(label)
    raw=openpyxl.load_workbook(ROOT/'data/raw/结果提交模板.xlsx',read_only=True)
    for sheet in raw:
        header=[x for x in next(sheet.values) if x is not None]
        assert list(next(w[sheet.title].values))[:len(header)]==header
    t=frame('transport_sorties');b=frame('box_delivery');r=frame('relay_sorties');g=frame('communication_guarantee')
    equal(rows('Q2_运输架次',8),rows_q2(t),'Q2_transport_actual_F13')
    equal(rows('Q3_运输时刻_必读'),rows('Q2_运输架次',8),'Same_Q2_Q3_transport_clock')
    equal(rows('Q2_逐箱交付',4),rows_delivery(b),'Q2_box_delivery_actual_F13')
    equal(rows('Q3_逐箱交付_必读'),rows('Q2_逐箱交付',4),'Same_Q2_Q3_delivery_clock')
    boxes=pd.read_csv(GEOMETRY/'boxes.csv').set_index('box_id')
    assert len(b)==80 and set(b.box_id)==set(boxes.index) and b.box_id.is_unique
    assert len(t)==24 and len(r)==3
    # Every supplement cell is re-read and checked against its execution source.
    for name,stem in [('运输完整时刻','transport_sorties'),('运输机资源日历','transport_uav_calendar'),
      ('运输电池资源日历','transport_battery_calendar'),('中继完整时刻','relay_sorties'),
      ('中继机资源日历','relay_uav_calendar'),('中继能源资源日历','relay_energy_calendar'),('原子通信保障关系','communication_guarantee')]:
        f=frame(stem)
        assert list(next(w[name].values))==list(f.columns)
        equal(rows(name),[[clean(x) for x in row] for row in f.itertuples(index=False,name=None)],name)
    # Verify official relay columns from full supplementary execution, including CRS.
    sites=frame('candidate_sites').set_index('site_id');tf=Transformer.from_crs(32649,4326,always_xy=True);expected=[]
    for x in r.itertuples():
        s=sites.loc[x.site_id];lon,lat=tf.transform(s.x_m,s.y_m)
        expected.append([x.relay_sortie_id,x.relay_id,x.energy_component_id,x.preparation_start_s,lon,lat,s.z_amsl_m,x.service_start_s,x.service_end_s,x.return_s,x.total_energy_kwh])
    equal(rows('Q3_中继架次',11),expected,'Official_relay_rows')
    comm=rows('Q3_通信保障',6)
    for f in t.itertuples():
        cc=sorted([z for z in comm if z[0]==f.sortie_id],key=lambda z:z[2])
        assert abs(cc[0][2]-f.takeoff_s)<1e-8 and abs(cc[-1][3]-f.return_s)<1e-8
        assert all(abs(a[3]-z[2])<1e-8 for a,z in zip(cc,cc[1:]))
        for z in cc:
            assert z[3]>z[2]
            active=g[(g.transport_sortie_id==f.sortie_id)&(g.service_start_s<(z[2]+z[3])/2)&(g.service_end_s>(z[2]+z[3])/2)]
            assert (len(active)==1 and z[4]=='TRANSPORT_RELAY_G01' and z[5]==active.iloc[0].relay_sortie_id) or (len(active)==0 and z[4]=='TRANSPORT_G01' and z[5] is None)
    checks.append('Whole_flight_communication_interval_partition')
    delivery={z[0]:z[3] for z in rows('Q2_逐箱交付',4)}
    late=sum(x.priority_weight*max(0,delivery[bid]-x.expected_deadline_s) for bid,x in boxes.iterrows())
    norm=sum(x.priority_weight*delivery[bid]/x.expected_deadline_s for bid,x in boxes.iterrows())/boxes.priority_weight.sum()
    m=read(PACKAGE/'joint_metrics_F13.json')
    assert abs(late-m['J_late'])<1e-8 and abs(norm-m['J_norm'])<1e-10
    indicators={x[0]:x[1] for x in rows('主指标核对')}
    energies=sum(z[7] for z in rows('Q2_运输架次',8))+sum(z[10] for z in rows('Q3_中继架次',11))
    makespan=max([z[6] for z in rows('Q2_运输架次',8)]+[z[9] for z in rows('Q3_中继架次',11)])
    assert abs(energies-m['total_energy_kwh'])<1e-8 and abs(makespan-m['joint_makespan_s'])<1e-8
    assert abs(indicators['Q3总能耗']-energies)<1e-8 and abs(indicators['Q3联合完工']-makespan)<1e-8
    assert abs(indicators['J_norm']-norm)<1e-10 and indicators['J_late']==late==0
    assert abs(indicators['Q3联合完工分钟']-makespan/60)<1e-9
    assert indicators['Q2/Q3运输架次']==24 and indicators['Q3中继架次']==3 and indicators['Gamma_C']==0
    assert abs(indicators['Q2/Q3运输能耗']-t.energy_kwh.sum())<1e-8
    assert abs(indicators['Q2/Q3运输完工']-t.return_s.max())<1e-8
    assert abs(indicators['Q3中继能耗']-r.total_energy_kwh.sum())<1e-8
    checks.append('Main_metrics_recomputed_from_workbook')
    q1=rows('Q1_单点组批',9);allboxes=[v for z in q1 for v in z[3].split(',')]
    assert len(allboxes)==80 and set(allboxes)==set(boxes.index)
    pk=pd.read_csv(ROOT/'results/reset/q1/q1_packings_rho20.csv',float_precision='round_trip')
    typ=pd.read_csv(GEOMETRY/'transport_uav_types.csv').set_index('uav_type')
    exp=[[f'Q1_{i:03d}',x.service_id,x.uav_type,','.join(ast.literal_eval(x.box_ids)),x.mass_kg,x.volume_m3,x.time_s,x.energy_kwh,100*(1-x.energy_kwh/typ.loc[x.uav_type].battery_energy_kwh)] for i,x in enumerate(pk.itertuples(),1)]
    equal(q1,exp,'G2_Q1_rows_and_SOC')
    assert indicators['Q1箱数']==80 and indicators['Q1架次']==len(pk) and abs(indicators['Q1能耗']-pk.energy_kwh.sum())<1e-8
    cfg=read(OUT/'q4/config.json');expected=[];full=[];deficits=[];work=[];service=[];calendar=[]
    for sid in dict.fromkeys(cfg['selected']+cfg['balanced_alternatives']):
        q=read(OUT/'q4/solutions'/sid/'configuration.json');role='缺口优先' if sid in cfg['selected'] else '均衡优先'
        for group in q['groups']:
            line=[q['k'],group['group_id'],','.join(group['services'])]+[group['resources'][k] for k in KINDS]
            if sid in cfg['selected']:expected.append(line)
            full.append([sid,role]+line)
            for svc in group['services']:service.append([sid,role,q['k'],svc,group['group_id']])
            work.append([sid,role,q['k'],group['group_id'],group['box_count'],group['mass_kg'],group['transport_sorties'],group['relay_sorties'],group['workload_uav_occupied_s'],q['workload_cv'],group['occupied_resource_time_s'],group['allocated_resource_time_s']])
        for k in KINDS:deficits.append([sid,role,q['k'],k,q['resources'][k],q['inventory'][k],q['shortage'][k],q['inventory_surplus'][k],q['redundancy'][k]])
        for x in q['calendar']:calendar.append([sid,role]+list(x.values()))
    equal(rows('Q4_分区配置',11),expected,'Q4_official_fixed_F13')
    equal(rows('Q4_双准则代表'),full,'Q4_both_selection_rules')
    equal(rows('Q4_分类型缺口冗余'),deficits,'Q4_deficits_surplus_redundancy')
    equal(rows('Q4_工作量与均衡'),work,'Q4_workload_CV')
    equal(rows('Q4_分区归属'),service,'Q4_service_assignment')
    equal(rows('Q4_独立资源日历'),calendar,'Q4_independent_calendars')
    assert not any(c.data_type=='f' for s in w for row in s for c in row)
    result=dict(status='PASSED',workbook=str(BOOK.relative_to(ROOT)),sha256=digest(BOOK),source_solution=IDENT,
      Q2_and_Q3_actual_clock_identical=True,Q4_source_identical=True,box_count=80,transport_sorties=24,relay_sorties=3,
      joint_makespan_s=makespan,total_energy_kwh=energies,J_late=late,J_norm=norm,
      checked_sheets=len(w.sheetnames),sheet_row_counts={s.title:s.max_row-1 for s in w},checks=checks,
      units='s,kWh,degrees,m; Q1 SOC percent; normalized index dimensionless',
      formal_main_promotion='CURRENT_AUTONOMOUS_MAIN_CANDIDATE; main branch unchanged')
    write(OUT/'submission_consistency.json',result);print('SUBMISSION_CHECK',result['checked_sheets'],checks,flush=True)
if __name__=='__main__':
    install_guard();build();check()

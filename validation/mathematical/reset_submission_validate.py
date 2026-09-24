"""Independent workbook/source equality and complete communication coverage."""
import ast,hashlib,json,sys
from pathlib import Path
import pandas as pd
import openpyxl
from pyproj import Transformer
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'results/reset';DEST=OUT/'submission'
def read(p):return pd.read_csv(p,float_precision='round_trip')
def close(a,b):
    if isinstance(b,(int,float)):assert abs(float(a)-float(b))<1e-7,(a,b)
    else:assert a==b,(a,b)
def validate():
    p=DEST/'结果提交表_G2_自主主方案.xlsx';w=openpyxl.load_workbook(p,data_only=True);raw=openpyxl.load_workbook(ROOT/'data/raw/结果提交模板.xlsx',data_only=True)
    manifest=json.loads((DEST/'input_manifest.json').read_text())
    for name,h in manifest['hashes'].items():assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==h,name
    widths={'Q1_单点组批':9,'Q2_运输架次':8,'Q2_逐箱交付':4,'Q3_中继架次':11,'Q3_通信保障':6,'Q4_分区配置':11}
    for name,n in widths.items():assert [w[name].cell(1,i).value for i in range(1,n+1)]==[raw[name].cell(1,i).value for i in range(1,n+1)]
    def compare(name,expected):
        actual=[r for r in w[name].iter_rows(min_row=2,max_col=len(expected[0]),values_only=True) if any(v is not None for v in r)]
        assert len(actual)==len(expected),(name,len(actual),len(expected))
        for a,b in zip(actual,expected):
            for x,y in zip(a,b):close(x,y)
    base=OUT/'q2/pareto_schedules/A11';q3=OUT/'q3/A11/solutions/A11_Q3_001'
    for directory,sheet_t,sheet_b,ft,fb in [(base,'Q2_运输架次','Q2_逐箱交付','q2_sorties.csv','q2_box_delivery.csv'),(q3,'Q3_运输时刻_必读','Q3_逐箱交付_必读','transport_sorties_A11.csv','box_delivery_A11.csv')]:
        compare(sheet_t,[[r.sortie_id,r.uav_id,r.uav_type,r.battery_id,r.preparation_start_s,r.service_sequence,r.return_s,r.energy_kwh] for r in read(directory/ft).itertuples()])
        compare(sheet_b,[[r.box_id,r.sortie_id,r.service_id,r.delivery_time_s] for r in read(directory/fb).itertuples()])
    types=read(OUT/'geometry/transport_uav_types.csv').set_index('uav_type')
    compare('Q1_单点组批',[[f'Q1_{i:03d}',r.service_id,r.uav_type,','.join(ast.literal_eval(r.box_ids)),r.mass_kg,r.volume_m3,r.time_s,r.energy_kwh,100*(1-r.energy_kwh/types.loc[r.uav_type].battery_energy_kwh)] for i,r in enumerate(read(OUT/'q1/q1_packings_rho20.csv').itertuples(),1)])
    sites=read(q3/'candidate_sites_A11.csv').set_index('site_id');tf=Transformer.from_crs(32649,4326,always_xy=True);expected=[]
    for r in read(q3/'relay_sorties_A11.csv').itertuples():
        s=sites.loc[r.site_id];lon,lat=tf.transform(s.x_m,s.y_m)
        expected.append([r.relay_sortie_id,r.relay_id,r.energy_component_id,r.preparation_start_s,lon,lat,s.z_amsl_m,r.service_start_s,r.service_end_s,r.return_s,r.total_energy_kwh])
    compare('Q3_中继架次',expected)
    cfg=json.loads((OUT/'q4/config.json').read_text());expected=[];kinds=['TUAV_A','TUAV_B','TUAV_C','TBAT_A','TBAT_B','TBAT_C','RUAV','REC']
    for sid in cfg['selected']:
        q=json.loads((OUT/'q4/solutions'/sid/'configuration.json').read_text())
        expected += [[q['k'],g['group_id'],','.join(g['services'])]+[g['resources'][k] for k in kinds] for g in q['groups']]
    compare('Q4_分区配置',expected)
    segments=[r for r in w['Q3_通信保障'].iter_rows(min_row=2,max_col=6,values_only=True) if r[0] is not None]
    guarantees=read(q3/'communication_guarantee_A11.csv');flights=read(q3/'transport_sorties_A11.csv')
    assert {x[0] for x in segments}==set(flights.sortie_id)
    for f in flights.itertuples():
        intervals=sorted((x for x in segments if x[0]==f.sortie_id),key=lambda x:x[2]);end=f.takeoff_s;gg=guarantees[guarantees.transport_sortie_id==f.sortie_id]
        for sid,phase,a,b,mode,rid in intervals:
            close(a,end);assert b>a;end=b
            overlaps=gg[(gg.service_start_s<b-1e-7)&(gg.service_end_s>a+1e-7)]
            if mode=='TRANSPORT_G01':assert overlaps.empty and rid is None
            else:
                assert mode=='TRANSPORT_RELAY_G01' and len(overlaps)==1
                r=overlaps.iloc[0];assert r.relay_sortie_id==rid and r.service_start_s<=a+1e-7 and r.service_end_s>=b-1e-7
        close(end,f.return_s)
    audit=json.loads((q3/'validation_0.25.json').read_text());assert audit['status']=='G2_Q3_INDEPENDENTLY_VALIDATED' and audit['uncovered_sample_count']==0
    result=dict(status='G2_SUBMISSION_INDEPENDENTLY_VERIFIED',official_sheets=6,Q1_sorties=18,Q2_transport_sorties=25,Q3_transport_sorties=25,Q3_relay_sorties=5,Q3_communication_intervals=len(segments),Q4_configurations=2,Q4_rows=5,explicit_Q3_transport_and_delivery_supplements=True,source_hashes_verified=True,workbook_sha256=hashlib.sha256(p.read_bytes()).hexdigest())
    (DEST/'validation.json').write_text(json.dumps(result,indent=2)+'\n');print(result)
if __name__=='__main__':validate()

#!/usr/bin/env python3
"""E0 Python smoke test. Does not read or modify contest raw data."""
from __future__ import annotations

import csv
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import pyproj
import rasterio
import scipy
from openpyxl import Workbook


def main() -> None:
    root = Path("/home/ccc/projects/uav-logistics-comms")
    out_dir = root / "results" / "e0"
    out_dir.mkdir(parents=True, exist_ok=True)

    # imports already done at module level
    versions = {
        "numpy": np.__version__,
        "scipy": scipy.__version__,
        "pandas": pd.__version__,
        "openpyxl": openpyxl.__version__,
        "matplotlib": matplotlib.__version__,
        "rasterio": rasterio.__version__,
        "pyproj": pyproj.__version__,
    }

    # temp Excel (NOT contest data)
    xlsx_path = out_dir / "python_smoke_tmp.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "smoke"
    ws.append(["k", "v"])
    ws.append(["status", "ok"])
    ws.append(["numpy", versions["numpy"]])
    wb.save(xlsx_path)

    # read back excel
    wb2 = openpyxl.load_workbook(xlsx_path)
    rows = list(wb2.active.iter_rows(values_only=True))
    assert rows[0] == ("k", "v")
    assert rows[1][0] == "status"

    # CSV write/read
    csv_path = out_dir / "python_smoke.csv"
    matrix = np.arange(12, dtype=float).reshape(3, 4)
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["lib", "version", "check"])
        for k, v in versions.items():
            w.writerow([k, v, "import_ok"])
        w.writerow(["matrix_rows", matrix.shape[0], "np_ok"])
        w.writerow(["matrix_cols", matrix.shape[1], "np_ok"])
    df = pd.read_csv(csv_path)
    assert len(df) >= 7

    # figure
    png_path = out_dir / "python_smoke.png"
    fig, ax = plt.subplots(figsize=(5, 3))
    ax.plot(matrix.mean(axis=1), marker="o")
    ax.set_title("E0 Python smoke")
    ax.set_xlabel("row")
    ax.set_ylabel("mean")
    fig.tight_layout()
    fig.savefig(png_path, dpi=120)
    plt.close(fig)

    print("python_smoke_ok")
    for k, v in versions.items():
        print(f"{k}={v}")
    print(f"wrote {csv_path}")
    print(f"wrote {png_path}")
    print(f"wrote {xlsx_path}")


if __name__ == "__main__":
    main()

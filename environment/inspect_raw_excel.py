#!/usr/bin/env python3
"""Inspect raw Excel/DEM structure for E1 ingestion design."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path("/home/ccc/projects/uav-logistics-comms")
RAW = ROOT / "data" / "raw"
BASE = RAW / "数据" / "无人机应急物资运输基础数据"

files = [
    BASE / "调度中心与服务区.xlsx",
    BASE / "运输无人机数据.xlsx",
    BASE / "中继无人机数据.xlsx",
    BASE / "物资需求与配送时限.xlsx",
    BASE / "通信链路参数.xlsx",
]

for f in files:
    print("=" * 70)
    print(f.name)
    wb = openpyxl.load_workbook(f, data_only=True)
    for ws in wb.worksheets:
        print(f"-- sheet: {ws.title} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i > 25:
                print("  ...")
                break
            print(" ", row)

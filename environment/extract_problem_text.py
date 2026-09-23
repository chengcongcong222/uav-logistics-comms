#!/usr/bin/env python3
"""Extract problem-statement text (esp. formulas) from the source document."""
from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path("/home/ccc/projects/uav-logistics-comms")
DOCX = ROOT / "data/raw/山区洪涝灾害下无人机运输与通信协同优化.docx"
OUT = ROOT / "results/e1/problem_statement_extract.txt"

doc = Document(DOCX)
parts = []
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        parts.append(t)
for ti, table in enumerate(doc.tables):
    parts.append(f"\n===== TABLE {ti} =====")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        parts.append(" | ".join(cells))

text = "\n".join(parts)
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(text, encoding="utf-8")
print(f"wrote {OUT} chars={len(text)}")
# print formula-ish regions
for i, line in enumerate(text.splitlines()):
    if any(k in line for k in ["能耗", "航程", "充电", "爬升", "下降", "SOC", "电量", "时间", "公式", "等效", "储备", "返航"]):
        print(f"{i}: {line}")
